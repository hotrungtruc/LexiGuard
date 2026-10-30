"""
Unit tests for LexiGuard DOCX Exporter Engine
"""

import io
import docx
from src.reporting.docx_exporter import generate_contract_audit_docx


def test_generate_contract_audit_docx_basic():
    report_data = {
        "health_score": 75,
        "executive_summary": "Tóm tắt rủi ro hợp đồng mẫu.",
        "tand_viac_summary": "Thực tiễn xét xử TAND và trọng tài VIAC.",
        "findings": [
            {
                "clause_reference": "Điều 5",
                "severity": "CRITICAL",
                "audit_verdict": "Vi phạm trần phạt 8%.",
                "statutory_citation": "Điều 301 LTM 2005"
            }
        ],
        "redline_recommendations": [
            {
                "clause_id": "Điều 5",
                "original_text": "Phạt 20%.",
                "balanced_revision": "Phạt tối đa 8%.",
                "strict_revision": "Phạt tối đa 5%.",
                "flexible_revision": "Phạt 8% có điều kiện.",
                "precedent_reference": "Án lệ 01",
                "rationale": "Tuân thủ luật thương mại."
            }
        ]
    }

    stream = generate_contract_audit_docx(report_data, contract_title="Hợp đồng thử nghiệm")
    assert isinstance(stream, io.BytesIO)
    data = stream.getvalue()
    assert len(data) > 5000

    doc = docx.Document(stream)
    assert len(doc.paragraphs) > 5
    assert len(doc.tables) >= 3


def test_generate_contract_audit_docx_empty():
    report_data = {
        "health_score": 95,
        "executive_summary": "Hợp đồng an toàn.",
        "findings": [],
        "redline_recommendations": []
    }

    stream = generate_contract_audit_docx(report_data, contract_title="Hợp đồng sạch")
    doc = docx.Document(stream)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "BÁO CÁO THẨM ĐỊNH RỦI RO HỢP ĐỒNG" in text

